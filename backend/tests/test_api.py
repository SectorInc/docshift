import io
import shutil
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from PIL import Image

from app.main import app, limiter

client = TestClient(app)
needs_gs = pytest.mark.skipif(not shutil.which("gs"), reason="Ghostscript not installed")
needs_lo = pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")


@pytest.fixture(autouse=True)
def reset_limiter():
    limiter.hits.clear()


def png(color="red") -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (300, 200), color).save(buf, "PNG")
    return buf.getvalue()


def docx_bytes() -> bytes:
    d = Document()
    d.add_heading("Hello", 1)
    d.add_paragraph("A short paragraph for testing the converter.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_pdf(n_images=1) -> bytes:
    files = [("files", (f"p{i}.png", png(), "image/png")) for i in range(n_images)]
    return client.post("/api/images-to-pdf", files=files).content


def upload(name, data, mime="application/pdf"):
    return ("files", (name, data, mime))


def test_health_and_limits():
    assert client.get("/api/health").json() == {"status": "ok"}
    lim = client.get("/api/limits").json()
    assert lim["max_files"] == 5 and lim["max_merge_files"] == 20


def test_images_to_pdf_combined():
    r = client.post("/api/images-to-pdf", files=[upload(f"{i}.png", png(), "image/png") for i in range(3)])
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


def test_images_to_pdf_separate_returns_zip():
    r = client.post("/api/images-to-pdf", data={"mode": "separate"},
                    files=[upload(f"{i}.png", png(), "image/png") for i in range(2)])
    assert r.headers["content-type"] == "application/zip"
    assert len(zipfile.ZipFile(io.BytesIO(r.content)).namelist()) == 2


def test_limit_of_five_files():
    r = client.post("/api/images-to-pdf", files=[upload(f"{i}.png", png(), "image/png") for i in range(6)])
    assert r.status_code == 400 and "5 files" in r.json()["detail"]


def test_merge_allows_twenty_and_rejects_twenty_one():
    pdf = make_pdf()
    ok = client.post("/api/merge", files=[upload(f"{i}.pdf", pdf) for i in range(20)])
    assert ok.status_code == 200 and ok.content.startswith(b"%PDF")
    too_many = client.post("/api/merge", files=[upload(f"{i}.pdf", pdf) for i in range(21)])
    assert too_many.status_code == 400


def test_merge_needs_two_files_and_counts_pages():
    pdf = make_pdf()
    assert client.post("/api/merge", files=[upload("a.pdf", pdf)]).status_code == 400
    from pypdf import PdfReader
    r = client.post("/api/merge", files=[upload("a.pdf", make_pdf(2)), upload("b.pdf", make_pdf(3))])
    assert len(PdfReader(io.BytesIO(r.content)).pages) == 5


def test_rejects_wrong_content():
    r = client.post("/api/merge", files=[upload("a.pdf", b"not a pdf"), upload("b.pdf", b"nope")])
    assert r.status_code == 400 and "not a valid .pdf" in r.json()["detail"]


def test_rejects_wrong_extension():
    r = client.post("/api/word-to-pdf", files=[upload("a.pdf", make_pdf())])
    assert r.status_code == 400


def test_rate_limit():
    for _ in range(30):
        client.post("/api/merge", files=[upload("a.pdf", b"x")])
    assert client.post("/api/merge", files=[upload("a.pdf", b"x")]).status_code == 429


def test_pdf_to_word():
    r = client.post("/api/pdf-to-word", files=[upload("a.pdf", make_pdf())])
    assert r.status_code == 200 and r.content.startswith(b"PK")


@needs_lo
def test_word_to_pdf():
    r = client.post("/api/word-to-pdf", files=[upload("a.docx", docx_bytes(), "application/octet-stream")])
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


@needs_gs
def test_compress_never_grows_file():
    pdf = make_pdf(3)
    r = client.post("/api/compress", files=[upload("a.pdf", pdf)])
    assert r.status_code == 200 and len(r.content) <= len(pdf)
